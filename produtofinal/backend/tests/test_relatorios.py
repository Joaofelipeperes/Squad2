"""Tela Relatórios (US20, PBI-60/61/77): catálogo, prévia e exportação XLSX/CSV. Dados fictícios."""
import io
import re
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from app.core.db import SessionLocal
from app.modules.acesso.service import criar_usuario
from app.modules.lgpd.models import Achado
from app.modules.relatorios import regras
from tests.conftest import SENHA, login
from tests.test_painel_organizacoes import _hoje, inventario  # noqa: F401 (fixtures)
from tests.test_painel_visao_geral import PDA, so_dicionario  # noqa: F401 (fixtures)

API = "/api/v1/relatorios"
IDS = ["pda", "atualizacoes", "lgpd", "orgaos", "sem-recurso", "formatos"]


@pytest.fixture
def com_pda(client, gerente, so_dicionario):  # noqa: F811
    r = client.post("/api/v1/pda/planos", headers=gerente, data={"nome": "PDA teste"},
                    files={"arquivo": ("pda.csv", PDA.encode(), "text/csv")})
    assert r.status_code == 201, r.text


@pytest.fixture
def achados(inventario):  # noqa: F811
    with SessionLocal() as db:
        db.add_all([
            Achado(recurso_id=1, coluna="cpf_paciente", linha=2, ocorrencias=10, tipo="cpf",
                   metodo="regra", confianca=95, prioridade="Alta", impressao_digital="a1"),
            Achado(recurso_id=1, coluna="email", linha=3, ocorrencias=1, tipo="email",
                   metodo="regra", confianca=80, prioridade="Média", impressao_digital="a2",
                   status="falso_positivo"),
        ])
        db.commit()


def _previa(client, headers, nome):
    r = client.get(f"{API}/{nome}/previa", headers=headers)
    assert r.status_code == 200, r.text
    corpo = r.json()
    return corpo, [dict(zip(corpo["colunas"], linha, strict=True)) for linha in corpo["linhas"]]


def _usuario(client, email, papeis):
    with SessionLocal() as db:
        criar_usuario(db, email, "Teste", SENHA, papeis)
    return login(client, email)


def test_catalogo_segue_permissao_de_origem(client, gerente):
    assert [r["id"] for r in client.get(API, headers=gerente).json()] == IDS
    # Superintendência só acessa o painel: só o relatório por órgão
    sup = _usuario(client, "sup@cge.go.gov.br", ["superintendencia"])
    assert [r["id"] for r in client.get(API, headers=sup).json()] == ["orgaos"]
    assert client.get(f"{API}/pda/previa", headers=sup).status_code == 403


def test_permissoes(client, orgao, gerente):
    assert client.get(API, headers=orgao).status_code == 403
    assert client.get(f"{API}/inexistente/previa", headers=gerente).status_code == 404
    assert client.get(f"{API}/pda/exportar?formato=doc", headers=gerente).status_code == 422


def test_pda_e_atualizacoes(client, gerente, com_pda):
    corpo, linhas = _previa(client, gerente, "pda")
    assert corpo["referencia"] == "PDA teste" and corpo["total"] == 6
    leitos = next(x for x in linhas if x["Base prevista"] == "Leitos")
    assert leitos["Situação"] == "Próximo do prazo" and leitos["Prazo de abertura"] == "31/10/2026"
    # Atualizações: só a vinculada atrasada (Indicadores, 01/09 + 30 dias)
    _, linhas = _previa(client, gerente, "atualizacoes")
    assert [(x["Base prevista"], x["Situação"], x["Vencimento"]) for x in linhas] == [
        ("Indicadores", "Atrasado", "01/10/2026")]


def test_lgpd_so_pendentes_de_revisao(client, gerente, achados):
    corpo, linhas = _previa(client, gerente, "lgpd")
    assert corpo["total"] == 1
    assert linhas[0]["Tipo"] == "CPF" and linhas[0]["Coluna"] == "cpf_paciente"
    assert linhas[0]["Prioridade"] == "Alta" and linhas[0]["Método"] == "Regra"


def test_inventario(client, gerente, so_dicionario):  # noqa: F811
    _, linhas = _previa(client, gerente, "sem-recurso")
    assert [x["Dataset"] for x in linhas] == ["Escolas"]
    _, linhas = _previa(client, gerente, "formatos")
    assert sorted(x["Formato"] for x in linhas) == ["(não informado)", "PDF", "ZIP"]


def test_exportar_xlsx(client, gerente, com_pda):
    r = client.get(f"{API}/pda/exportar?formato=xlsx", headers=gerente)
    assert r.status_code == 200
    assert 'filename="relatorio-pda-2026-10-18.xlsx"' in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "Sigla" and ws["A1"].font.bold
    assert ws.max_row == 7  # cabeçalho + 6 bases


def test_exportar_csv(client, gerente, com_pda):
    r = client.get(f"{API}/orgaos/exportar?formato=csv", headers=gerente)
    assert r.status_code == 200
    texto = r.content.decode("utf-8")
    assert texto.startswith("﻿Órgão;Siglas no PDA;")
    assert "Secretaria da Saúde;SES;4;" in texto


@pytest.mark.parametrize("valor,esperado", [
    (27.0, "27"), (33.3, "33,3"), (None, ""), (True, "Sim"), (date(2026, 10, 1), "01/10/2026"),
])
def test_texto_da_celula(valor, esperado):
    assert regras.texto(valor) == esperado


def test_exportar_pdf(client, gerente, com_pda):
    r = client.get(f"{API}/pda/exportar?formato=pdf", headers=gerente)
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/pdf"
    assert 'filename="relatorio-pda-2026-10-18.pdf"' in r.headers["content-disposition"]
    assert r.content.startswith(b"%PDF") and b"/Title (Relat" in r.content


def test_pdf_muitas_linhas_e_vazio():
    """Tabela longa quebra em várias páginas; relatório vazio ainda gera o PDF."""
    agora = datetime(2026, 10, 30, 9, 0)
    linhas = [[f"Base {i} com nome comprido & <especial>", "Publicado", date(2026, 1, 1)]
              for i in range(300)]
    longo = regras.gerar_pdf("Teste", ["Base", "Situação", "Prazo"], linhas, referencia="PDA",
                             gerado_em=agora)
    assert int(re.search(rb"/Count (\d+)", longo).group(1)) > 1  # páginas no catálogo do PDF
    vazio = regras.gerar_pdf("Teste", ["Base", "Situação"], [], referencia=None, gerado_em=agora)
    assert vazio.startswith(b"%PDF")


def test_exportar_exige_permissao(client, admin, gerente):
    """Sem relatorios.exportar: vê a prévia, mas não exporta."""
    r = client.post("/api/v1/acesso/papeis", headers=admin, json={
        "codigo": "so_ver", "nome": "Só ver",
        "permissoes": ["relatorios.acessar", "painel.acessar"]})
    assert r.status_code == 201, r.text
    ver = _usuario(client, "ver@cge.go.gov.br", ["so_ver"])
    assert client.get(f"{API}/orgaos/previa", headers=ver).status_code == 200
    assert client.get(f"{API}/orgaos/exportar", headers=ver).status_code == 403
