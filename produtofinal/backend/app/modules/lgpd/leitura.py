"""Leitura de recursos tabulares para a varredura LGPD (PBI-32), sem pandas.

Entrega o arquivo como tabelas de texto: nome da tabela (aba do .xlsx; vazio no CSV), nomes das
colunas e blocos de linhas `[(numero_da_linha, [valores])]`. O número da linha é o do arquivo
(cabeçalho = linha 1, como no Excel), para a pessoa achar o dado.

- CSV/TSV/TXT: encoding detectado (utf-8 com ou sem BOM; senão latin-1) e separador detectado
  (`, ; tab |`), como no detector3.
- XLSX: todas as abas; a primeira linha não vazia de cada aba é o cabeçalho.
- Outros formatos (PDF, XLS antigo, JSON, ZIP...) → `FormatoNaoSuportado`.

Colunas sem nome ou repetidas ganham o número da posição: "escola (#3)".
Funções puras (sem banco, sem FastAPI).
"""
from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass

FORMATOS_CSV = {"CSV", "TSV", "TXT"}
FORMATOS_XLSX = {"XLSX"}
FORMATOS_SUPORTADOS = FORMATOS_CSV | FORMATOS_XLSX
LINHAS_POR_BLOCO = 50_000  # o perfil de coluna é calculado por bloco, como no detector3
DESCOMPACTADO_MAXIMO = 300 * 1024 * 1024  # proteção contra "bomba" de compressão no .xlsx


class FormatoNaoSuportado(ValueError):
    pass


class ArquivoInvalido(ValueError):
    pass


@dataclass
class Tabela:
    nome: str              # aba do .xlsx ("" no CSV)
    colunas: list[str]
    blocos: Iterator[list[tuple[int, list[str]]]]


def formato_do_recurso(formato: str | None, nome_ou_url: str | None) -> str | None:
    """Formato declarado no CKAN; se vazio, a extensão do nome/URL."""
    f = (formato or "").strip().upper().lstrip(".")
    f = f.removeprefix("TEXT/")  # alguns órgãos preenchem o tipo MIME ("text/csv")
    if f:
        return f
    nome = (nome_ou_url or "").split("?")[0].lower()
    for ext in ("csv", "tsv", "txt", "xlsx"):
        if nome.endswith("." + ext):
            return ext.upper()
    return None


def nomes_de_colunas(cabecalho: list[object]) -> list[str]:
    vistos: dict[str, int] = {}
    nomes = []
    for i, c in enumerate(cabecalho):
        nome = str(c).strip() if c is not None else ""
        vistos[nome] = vistos.get(nome, 0) + 1
        nomes.append(nome)
    return [n if n and vistos[n] == 1 else f"{n or 'coluna'} (#{i + 1})" for i, n in enumerate(nomes)]


def _decodificar(conteudo: bytes) -> str:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            return conteudo.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ArquivoInvalido("Não foi possível decodificar o arquivo.")  # latin-1 nunca falha


def _blocos(linhas: Iterator[tuple[int, list[str]]], tamanho: int):
    bloco: list[tuple[int, list[str]]] = []
    for item in linhas:
        bloco.append(item)
        if len(bloco) >= tamanho:
            yield bloco
            bloco = []
    if bloco:
        yield bloco


def ler_csv(conteudo: bytes, tamanho_bloco: int = LINHAS_POR_BLOCO) -> list[Tabela]:
    texto = _decodificar(conteudo)
    try:
        sep = csv.Sniffer().sniff(texto[:65536], delimiters=",;\t|").delimiter
    except csv.Error:
        sep = ","
    leitor = csv.reader(io.StringIO(texto, newline=""), delimiter=sep)
    try:
        cabecalho = next(leitor)
    except StopIteration:
        return []
    colunas = nomes_de_colunas(cabecalho)

    def linhas():
        for n, row in enumerate(leitor, start=2):
            if any(c.strip() for c in row):
                yield n, [(row[i] if i < len(row) else "") for i in range(len(colunas))]

    return [Tabela("", colunas, _blocos(linhas(), tamanho_bloco))]


def _texto_celula(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))  # 62998765432.0 → "62998765432" (números longos lidos como float)
    return str(v)


def ler_xlsx(conteudo: bytes, tamanho_bloco: int = LINHAS_POR_BLOCO) -> list[Tabela]:
    from openpyxl import load_workbook

    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
            if sum(i.file_size for i in z.infolist()) > DESCOMPACTADO_MAXIMO:
                raise ArquivoInvalido("Planilha grande demais depois de descompactada.")
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except (zipfile.BadZipFile, OSError, KeyError) as exc:
        raise ArquivoInvalido(f"Arquivo .xlsx ilegível: {exc}") from exc
    tabelas = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        cabecalho, n = None, 0
        for n, row in enumerate(it, start=1):
            if any(v not in (None, "") for v in row):
                cabecalho = list(row)
                break
        if cabecalho is None:
            continue
        colunas = nomes_de_colunas(cabecalho)

        def linhas(it=it, inicio=n + 1, largura=len(colunas)):
            for k, row in enumerate(it, start=inicio):
                valores = [_texto_celula(row[i] if i < len(row) else None) for i in range(largura)]
                if any(v.strip() for v in valores):
                    yield k, valores

        nome = ws.title if len(wb.worksheets) > 1 else ""
        tabelas.append(Tabela(nome, colunas, _blocos(linhas(), tamanho_bloco)))
    return tabelas


_ASSINATURAS = {b"PK\x03\x04": "compactado (ZIP)", b"%PDF": "PDF", b"\x1f\x8b": "compactado (GZIP)",
                b"Rar!": "compactado (RAR)", b"7z\xbc\xaf": "compactado (7z)",
                b"\xd0\xcf\x11\xe0": "planilha antiga do Excel (XLS)"}


def conteudo_real(conteudo: bytes) -> str | None:
    """O que o arquivo é de fato, pelos primeiros bytes, quando não é texto (None = texto)."""
    for assinatura, nome in _ASSINATURAS.items():
        if conteudo.startswith(assinatura):
            return nome
    return None


def ler(conteudo: bytes, formato: str | None, tamanho_bloco: int = LINHAS_POR_BLOCO) -> list[Tabela]:
    f = (formato or "").upper()
    real = conteudo_real(conteudo)
    if f in FORMATOS_CSV and real:
        raise FormatoNaoSuportado(f"O recurso está declarado como {f}, mas o arquivo é {real}: "
                                  "ainda não é varrido.")
    if f in FORMATOS_CSV:
        return ler_csv(conteudo, tamanho_bloco)
    if f in FORMATOS_XLSX:
        return ler_xlsx(conteudo, tamanho_bloco)
    raise FormatoNaoSuportado(f"Formato {f or 'desconhecido'} ainda não é varrido "
                              f"(suportados: {', '.join(sorted(FORMATOS_SUPORTADOS))}).")
