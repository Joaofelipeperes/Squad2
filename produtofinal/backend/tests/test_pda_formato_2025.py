"""Formato da planilha do PDA exportada pela GEDA (PDA 2025-2027): colunas Orgão, Base de Dados,
Descrição, Unidade Responsável, Atualização e "Meta/Prazo para abertura" (mês/ano), sem
"Disponível no Portal", com o rodapé "Filtros aplicados:Ano é 2025". Linhas copiadas da planilha
pública do PDA (sem dados pessoais)."""
import io
from datetime import date

import pytest
from openpyxl import Workbook

from app.modules.pda import regras, service
from tests.test_pda_revisao import _importar

API = "/api/v1/pda"
CAB = "Orgão;Base de Dados;Descrição;Unidade Responsável;Atualização;Meta/Prazo para abertura"
LINHAS = [
    "VICE GOVERNADORIA;Bens Móveis;Lista dos Bens Móveis;Vice Governadoria;Mensal;Março/2025",
    "SEMAD;Embargos Estaduais polígono - INÃ/SGA;Poligonais dos Embargos;SUF;Mensal;Dezembro2025",
    ("DGPC;Atribuição das DEAMs;Atribuições das Delegacias;Autoridade de monitoramento;Semestral;"
     "Setembro 2025"),
    ("GOIÁSFOMENTO;Relação de Terceirizados;Relação de Terceirizados;GEPAT;"
     "Mensal/Trimestral/Semestral;Dezembro/2025"),
    "SES;Vacinas/Doses distribuídas;Doses distribuídas;Gerência de Imunização;N/A;Dezembro/2025",
]
CSV = "\n".join([CAB, *LINHAS, "", "Filtros aplicados:Ano é 2025"]) + "\n"


@pytest.fixture(autouse=True)
def _hoje(monkeypatch):
    monkeypatch.setattr(service, "hoje_local", lambda: date(2026, 10, 8))


@pytest.mark.parametrize("texto,esperado", [
    ("Março/2025", date(2025, 3, 31)), ("Dezembro2025", date(2025, 12, 31)),
    ("Setembro 2025", date(2025, 9, 30)), ("Fevereiro/2028", date(2028, 2, 29)),
    ("fev/2026", date(2026, 2, 28)), ("Junho de 2025", date(2025, 6, 30)),
    ("03/2025", date(2025, 3, 31)), ("ABRIL/2025", date(2025, 4, 30)),
])
def test_prazo_mes_e_ano_vale_ate_o_fim_do_mes(texto, esperado):
    assert regras.interpretar_data(texto) == esperado


@pytest.mark.parametrize("texto", ["13/2025", "Marçço/2025", "2025"])
def test_prazo_mes_e_ano_ilegivel(texto):
    with pytest.raises(ValueError):
        regras.interpretar_data(texto)


def test_le_o_formato_da_planilha_2025():
    lida = regras.ler_planilha(CSV.encode("utf-8"), "pda.csv")
    assert [(ln.orgao_sigla, ln.prazo) for ln in lida.linhas] == [
        ("VICE GOVERNADORIA", date(2025, 3, 31)), ("SEMAD", date(2025, 12, 31)),
        ("DGPC", date(2025, 9, 30)), ("GOIÁSFOMENTO", date(2025, 12, 31)),
        ("SES", date(2025, 12, 31))]
    assert lida.ignoradas == []  # o rodapé "Filtros aplicados" não é base nem erro
    assert lida.tem_coluna_prazo and not lida.tem_coluna_portal
    assert lida.colunas_opcionais_ausentes == [
        "Políticas Públicas", "Possui Conteúdo Sigiloso?", "Disponível no Portal"]
    assert [regras.normalizar_periodicidade(ln.periodicidade_original) for ln in lida.linhas] == [
        "Mensal", "Mensal", "Semestral", "Múltipla", "Sem informação"]


def test_xlsx_com_mes_e_ano_em_celula_de_data():
    """No Excel em pt-BR, "Março/2025" digitado vira a data 01/03/2025 com formato mmmm/aaaa."""
    wb = Workbook()
    ws = wb.active
    ws.append(CAB.split(";"))
    ws.append(["SES", "Dengue", "Dados", "SUVISA", "Mensal", date(2025, 9, 1)])
    ws["F2"].number_format = "mmmm/yyyy"
    ws.append(["AGR", "Gestão Tarifária", "Tarifas", "GERE", "Anual", date(2025, 5, 1)])
    ws["F3"].number_format = '[$-416]mmm\\-yy;@'
    ws.append(["CGE", "PAR e PAF", "Processos", "SCC", "Bimestral", date(2025, 6, 15)])
    ws["F4"].number_format = "dd/mm/yyyy"  # data completa: vale o dia
    ws.append(["UEG", "Obras", "Obras", "Escritório", "Mensal", "Junho/2025"])  # texto
    buf = io.BytesIO()
    wb.save(buf)
    lida = regras.ler_planilha(buf.getvalue(), "pda.xlsx")
    assert [ln.prazo for ln in lida.linhas] == [
        date(2025, 9, 30), date(2025, 5, 31), date(2025, 6, 15), date(2025, 6, 30)]


def test_importacao_do_formato_2025(client, gerente):
    r = _importar(client, gerente, "PDA 2025-2027", CSV)
    assert r.status_code == 201, r.text
    corpo = r.json()
    assert corpo["bases_importadas"] == 5 and corpo["linhas_ignoradas"] == []
    assert corpo["tem_coluna_prazo"] is True and corpo["tem_coluna_portal"] is False
    assert corpo["sem_vinculo"] == 5
    bases = client.get(f"{API}/bases", headers=gerente,
                       params={"plano_id": corpo["plano"]["id"]}).json()["bases"]
    # Sem vínculo e com a meta vencida (hoje = 08/10/2026) → Em atraso
    assert {b["situacao"] for b in bases} == {"Em atraso"}
    assert {b["prazo_abertura"] for b in bases} == {"2025-03-31", "2025-09-30", "2025-12-31"}
