"""Regras puras do módulo relatorios: catálogo dos relatórios e geração de CSV/XLSX/PDF (PBI-60/61).
Sem banco nem FastAPI (reaproveitáveis numa extensão CKAN — ADR-0004)."""
import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

FORMATOS = ("xlsx", "csv", "pdf")
MIME = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "csv": "text/csv; charset=utf-8", "pdf": "application/pdf"}
LARGURA_MAXIMA = 60  # colunas do XLSX (em caracteres)


@dataclass(frozen=True)
class Definicao:
    """Relatório do catálogo. `permissao_origem` = permissão de acesso do módulo de onde vêm os
    dados (além de relatorios.acessar)."""

    id: str
    titulo: str
    descricao: str
    icone: str
    permissao_origem: str


# Ordem e textos do protótipo V1 (#screen-reports, PBI-77)
CATALOGO: tuple[Definicao, ...] = (
    Definicao("pda", "Relatório de Cumprimento do PDA",
              "Situação das bases previstas no Plano de Dados Abertos vigente, por órgão e prazo.",
              "fa-clipboard-check", "pda.acessar"),
    Definicao("atualizacoes", "Relatório de Atualizações",
              "Bases vinculadas com atualização em atraso ou próxima do vencimento, pela "
              "periodicidade do PDA e o last_modified dos recursos.",
              "fa-rotate", "atualizacoes.acessar"),
    Definicao("lgpd", "Relatório de Dados Pessoais",
              "Sinalizações de possíveis dados pessoais pendentes de revisão humana, por "
              "prioridade.", "fa-shield-halved", "lgpd.acessar"),
    Definicao("orgaos", "Relatório por Órgão",
              "Resumo de cumprimento do PDA, atualização e formatos por órgão.",
              "fa-building-columns", "painel.acessar"),
    Definicao("sem-recurso", "Relatório de Datasets sem Recurso",
              "Datasets ativos no portal sem nenhum recurso válido publicado.",
              "fa-circle-question", "inventario.acessar"),
    Definicao("formatos", "Relatório de Formatos não Abertos",
              "Recursos publicados em formatos fora da lista de formatos abertos da GEDA.",
              "fa-file-circle-exclamation", "inventario.acessar"),
)
POR_ID = {d.id: d for d in CATALOGO}


def texto(valor) -> str:
    """Valor de célula como texto (CSV e prévia): datas em dd/mm/aaaa, vazio para None."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    if isinstance(valor, float):  # padrão brasileiro; inteiro sem ",0"
        return str(int(valor)) if valor.is_integer() else f"{valor:.1f}".replace(".", ",")
    return str(valor)


def gerar_csv(colunas: list[str], linhas: list[list]) -> bytes:
    """CSV para o Excel em pt-BR: BOM UTF-8 e ';' como separador (mesmo padrão do log de acesso)."""
    saida = io.StringIO()
    escritor = csv.writer(saida, delimiter=";", lineterminator="\r\n")
    escritor.writerow(colunas)
    escritor.writerows([texto(v) for v in linha] for linha in linhas)
    return ("﻿" + saida.getvalue()).encode("utf-8")


def _celula(valor):
    if isinstance(valor, datetime):
        return valor.replace(tzinfo=None)  # openpyxl não grava fuso horário
    if isinstance(valor, bool):
        return "Sim" if valor else "Não"
    return valor


def gerar_xlsx(titulo: str, colunas: list[str], linhas: list[list]) -> bytes:
    """Planilha com cabeçalho em negrito, cabeçalho congelado, filtro e largura ajustada."""
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]  # limite do Excel para nome de aba
    ws.append(colunas)
    for c in ws[1]:
        c.font = Font(bold=True)
    for linha in linhas:
        ws.append([_celula(v) for v in linha])
    for i, nome in enumerate(colunas, start=1):
        letra = get_column_letter(i)
        maior = max([len(nome), *(len(texto(linha[i - 1])) for linha in linhas)])
        ws.column_dimensions[letra].width = min(maior + 2, LARGURA_MAXIMA)
        for (c,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            if isinstance(c.value, datetime):
                c.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(c.value, date):
                c.number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A2"
    if linhas:
        ws.auto_filter.ref = ws.dimensions
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


PDF_FONTE = 7.5  # pt, células e cabeçalho
PDF_FOLGA = 8  # pt, padding esquerdo + direito da célula
PDF_LARGURA_DESEJADA = 220  # pt, teto do texto de uma célula numa linha só


def _larguras(colunas: list[str], linhas: list[list], total: float) -> list[float]:
    """Larguras somando `total`: cada coluna recebe ao menos a sua maior palavra (não quebra
    palavra no meio) e o espaço que sobra vai para as colunas de texto mais longo."""
    amostra = linhas[:300]
    minimas, desejadas = [], []
    for i, nome in enumerate(colunas):
        textos = [(nome, "Helvetica-Bold")] + [(texto(linha[i]), "Helvetica") for linha in amostra]
        palavra = max(stringWidth(p, fonte, PDF_FONTE) for t, fonte in textos for p in t.split()
                      or [""])
        inteiro = max(stringWidth(t, fonte, PDF_FONTE) for t, fonte in textos)
        minimas.append(palavra + PDF_FOLGA)
        desejadas.append(max(min(inteiro, PDF_LARGURA_DESEJADA) + PDF_FOLGA, minimas[-1]))
    if sum(minimas) >= total:  # nem as palavras cabem: proporcional às mínimas
        return [total * m / sum(minimas) for m in minimas]
    extra = [d - m for d, m in zip(desejadas, minimas, strict=True)]
    sobra = total - sum(minimas)
    if sum(extra) <= sobra:  # tudo cabe numa linha: distribui a folga pelo desejado
        return [total * d / sum(desejadas) for d in desejadas]
    return [m + sobra * e / sum(extra) for m, e in zip(minimas, extra, strict=True)]


def gerar_pdf(titulo: str, colunas: list[str], linhas: list[list], *, referencia: str | None,
              gerado_em: datetime) -> bytes:
    """PDF em A4 paisagem (ADR-0026): título, referência, data de geração e total; tabela com o
    cabeçalho repetido em cada página e texto quebrado nas células; rodapé com a paginação."""
    estilos = getSampleStyleSheet()
    celula = ParagraphStyle("celula", parent=estilos["BodyText"], fontSize=PDF_FONTE, leading=9)
    cabecalho = ParagraphStyle("cabecalho", parent=celula, fontName="Helvetica-Bold",
                               textColor=colors.HexColor("#334155"))
    pagina = landscape(A4)
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=pagina, leftMargin=1 * cm, rightMargin=1 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm, title=titulo,
                            author="Monitor Dados Abertos GO — GEDA/CGE-GO")

    subtitulo = " · ".join(filter(None, [
        referencia, f"Gerado em {gerado_em.strftime('%d/%m/%Y %H:%M')}",
        f"{len(linhas)} {'linha' if len(linhas) == 1 else 'linhas'}"]))
    corpo = [Paragraph(escape(titulo), estilos["Title"]),
             Paragraph(escape(subtitulo), estilos["Normal"]), Spacer(1, 0.4 * cm)]

    dados = [[Paragraph(escape(c), cabecalho) for c in colunas]]
    dados += [[Paragraph(escape(texto(v)), celula) for v in linha] for linha in linhas]
    if not linhas:
        dados.append([Paragraph("Nenhum registro.", celula)] + [""] * (len(colunas) - 1))
    tabela = LongTable(dados, colWidths=_larguras(colunas, linhas, doc.width), repeatRows=1)
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor("#CBD5E1")),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#E2E8F0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        *([("SPAN", (0, 1), (-1, 1))] if not linhas else []),
    ]))
    corpo.append(tabela)

    def rodape(canvas, documento):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(documento.leftMargin, 0.6 * cm, titulo)
        canvas.drawRightString(pagina[0] - documento.rightMargin, 0.6 * cm,
                               f"Página {documento.page}")
        canvas.restoreState()

    doc.build(corpo, onFirstPage=rodape, onLaterPages=rodape)
    return buffer.getvalue()


def nome_arquivo(relatorio_id: str, formato: str, hoje: date) -> str:
    return f"relatorio-{relatorio_id}-{hoje.isoformat()}.{formato}"
