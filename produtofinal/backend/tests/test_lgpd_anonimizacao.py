"""Anonimização (US26): operadores, reescrita e o fluxo prévia → publicação → reversão, sem rede."""
import io
from pathlib import Path

import httpx
import pytest
from openpyxl import Workbook, load_workbook
from pydantic import SecretStr
from test_lgpd import CPF_A, CPF_B, REC_CSV, _configurar_ia, csv_alunos, portal  # noqa: F401  (fixture)

from app.core.config import get_settings
from app.core.db import SessionLocal
from app.integrations.ckan.client import CkanClient
from app.integrations.ckan.writer import CkanWriter
from app.modules.inventario.models import Recurso
from app.modules.lgpd import anonimizacao, correcoes, deteccao
from app.modules.lgpd.models import Achado, CorrecaoPublicada, DecisaoTriagem


# ------------------------------------------------------------------ funções puras
def test_operadores():
    assert anonimizacao.mascara_cpf("529.982.247-25") == "***.982.247-**"   # regra da GEDA
    assert anonimizacao.mascara_cpf("52998224725") == "***982247**"
    assert anonimizacao.aplicar("substituir", "x", "email") == "<EMAIL>"
    assert anonimizacao.aplicar("mascarar", "aluno1@email.com", "email") == "alu***********om"
    h = anonimizacao.aplicar("hash", " 529.982.247-25", "cpf", "sal")
    assert h == anonimizacao.aplicar("hash", "529.982.247-25", "cpf", "sal") and len(h) == 12
    assert h != anonimizacao.aplicar("hash", "529.982.247-25", "cpf", "outro sal")
    assert anonimizacao.operador_padrao("cpf") == "mascara_cpf"
    assert anonimizacao.operador_padrao("telefone") == "substituir"


def test_planejar_trecho_detectado_e_coluna_inteira():
    analisador = deteccao.montar_analisador(0.5)
    blocos = [[(2, ["cel (62) 99876-5432", CPF_A]), (3, ["sem telefone", "123.456.789-00"])]]
    tel = anonimizacao.ItemPlano("1", "", 0, "telefone", "TELEFONE", "substituir", "detectados")
    cpf = anonimizacao.ItemPlano("2", "", 1, "cpf", "CPF", "mascara_cpf", "coluna")
    alt = anonimizacao.Alteracoes()
    anonimizacao.planejar({0: [tel], 1: [cpf]}, blocos,
                          lambda v, i: deteccao.analisar_coluna(analisador, v), "", alt, "")
    assert alt.celulas == {("", 2, 0): "cel <TELEFONE>",
                           ("", 2, 1): "***.982.247-**",
                           ("", 3, 1): "***.456.789-**"}  # coluna inteira pega até CPF inválido
    assert (tel.alteradas, cpf.alteradas, cpf.linhas_exemplo) == (1, 2, [2, 3])


def test_reescrita_preserva_formato():
    original = 'nome;obs\r\nAna;"tel 62 3241-5566; ramal"\r\n'.encode("latin-1")
    novo = anonimizacao.reescrever(original, "CSV",
                                   anonimizacao.Alteracoes({("", 2, 1): "tel <TELEFONE>; ramal"}))
    assert novo == 'nome;obs\r\nAna;tel <TELEFONE>; ramal\r\n'.encode("latin-1").replace(
        b"Ana;tel <TELEFONE>; ramal", b'Ana;"tel <TELEFONE>; ramal"')
    wb = Workbook()
    wb.active.append(["nome", "cpf"])
    wb.active.append(["Ana", CPF_A])
    buf = io.BytesIO()
    wb.save(buf)
    saida = anonimizacao.reescrever(buf.getvalue(), "XLSX",
                                    anonimizacao.Alteracoes({("", 2, 1): "***.982.247-**"}))
    ws = load_workbook(io.BytesIO(saida)).active
    assert [c.value for c in ws[2]] == ["Ana", "***.982.247-**"]


# ------------------------------------------------------------------ fluxo pela API
@pytest.fixture
def ambiente(portal, monkeypatch, tmp_path):  # noqa: F811
    """Arquivos fora do banco numa pasta temporária; CKAN simulado também para a correção."""
    monkeypatch.setattr(get_settings(), "lgpd_arquivos_dir", str(tmp_path))
    servido = {"conteudo": csv_alunos()}
    uploads: list[dict] = []

    def ckan(req: httpx.Request):
        if req.url.path.endswith("/resource_patch"):
            corpo = req.read()
            uploads.append({"corpo": corpo, "auth": req.headers.get("Authorization")})
            return httpx.Response(200, json={"success": True, "result": {"id": "rec-csv"}})
        return httpx.Response(200, content=servido["conteudo"])

    t = httpx.MockTransport(ckan)
    monkeypatch.setattr(correcoes, "CkanClient", lambda: CkanClient("http://ckan:5000", transport=t))
    monkeypatch.setattr(correcoes, "_writer", lambda: CkanWriter(transport=t))
    return {"servido": servido, "uploads": uploads, "pasta": tmp_path}


def _escrita(monkeypatch, ligada: bool):
    monkeypatch.setattr(get_settings(), "ckan_write_enabled", ligada)
    monkeypatch.setattr(get_settings(), "ckan_write_api_token", SecretStr("token-da-geda"))


def _varrer_e_confirmar(client, admin, gerente):
    _configurar_ia(client, admin)
    client.post("/api/v1/lgpd/varreduras", headers=gerente, json={"escopo": "recurso", "id": REC_CSV})
    achados = {a["tipo"]: a["id"] for a in
               client.get("/api/v1/lgpd/achados", headers=gerente).json()["achados"]}
    for tipo in ("cpf", "email"):
        client.post(f"/api/v1/lgpd/achados/{achados[tipo]}/decisao", headers=gerente,
                    json={"decisao": "confirmado"})
    return achados


def _previa(client, gerente, achados):
    r = client.post("/api/v1/lgpd/correcoes", headers=gerente, json={"recurso_id": REC_CSV, "itens": [
        {"achado_id": achados["cpf"], "operador": "mascara_cpf", "alcance": "detectados"},
        {"achado_id": achados["email"], "operador": "substituir", "alcance": "coluna"}]})
    assert r.status_code == 202, r.text
    return client.get(f"/api/v1/lgpd/correcoes/{r.json()['id']}", headers=gerente).json()


def test_previa_publicacao_e_reversao(client, admin, gerente, analista, ambiente, monkeypatch):
    achados = _varrer_e_confirmar(client, admin, gerente)
    c = _previa(client, gerente, achados)
    assert c["status"] == "previa", c["erro"]
    assert c["celulas_alteradas"] == 60 and c["nome_arquivo"] == "alunos.csv"
    itens = {i["tipo"]: i for i in c["itens"]}
    assert itens["cpf"]["alteradas"] == 30 and itens["email"]["alteradas"] == 30
    assert itens["cpf"]["exemplos"][0] == {"linha": 2, "valor": "***.982.247-**"}
    assert itens["email"]["exemplos"][0]["valor"] == "<EMAIL>"

    with SessionLocal() as db:
        cor = db.get(CorrecaoPublicada, c["id"])
        original, previa = Path(cor.caminho_backup), Path(cor.caminho_previa)
        assert "aluno1@email.com" not in str(cor.itens) and CPF_A not in str(cor.itens)
    assert original.read_bytes() == csv_alunos()
    assert oct(original.stat().st_mode)[-3:] == "600"
    texto = previa.read_bytes().decode("latin-1")
    assert CPF_A not in texto and CPF_B not in texto and "@email.com" not in texto
    assert texto.count("\n") == csv_alunos().decode("latin-1").count("\n")

    r = client.get(f"/api/v1/lgpd/correcoes/{c['id']}/arquivo-previa", headers=gerente)
    assert r.status_code == 200 and r.content == previa.read_bytes()
    assert client.get(f"/api/v1/lgpd/correcoes/{c['id']}/arquivo-previa",
                      headers=analista).status_code == 403

    # escrita desligada: a prévia fica, nada é publicado
    _escrita(monkeypatch, False)
    r = client.post(f"/api/v1/lgpd/correcoes/{c['id']}/aprovar", headers=gerente, json={})
    assert r.status_code == 409 and "desabilitada" in r.json()["detail"] and not ambiente["uploads"]

    _escrita(monkeypatch, True)
    r = client.post(f"/api/v1/lgpd/correcoes/{c['id']}/aprovar", headers=gerente,
                    json={"justificativa": "CPF e e-mail de alunos expostos"})
    assert r.status_code == 200 and r.json()["status"] == "publicada", r.text
    assert r.json()["aprovado_por"] == "paloma@cge.go.gov.br"
    (up,) = ambiente["uploads"]
    assert up["auth"] == "token-da-geda" and b'name="id"' in up["corpo"] and b"rec-csv" in up["corpo"]
    assert previa.read_bytes() in up["corpo"] and b'filename="alunos.csv"' in up["corpo"]
    with SessionLocal() as db:
        assert {a.status for a in db.query(Achado) if a.tipo in ("cpf", "email")} == {"tratado"}
        assert db.query(DecisaoTriagem).filter_by(decisao="tratado").count() == 2

    r = client.post(f"/api/v1/lgpd/correcoes/{c['id']}/reverter", headers=gerente,
                    json={"motivo": "Base com autorização de publicação"})
    assert r.status_code == 200 and r.json()["status"] == "revertida"
    assert csv_alunos() in ambiente["uploads"][1]["corpo"]
    with SessionLocal() as db:
        assert {a.status for a in db.query(Achado) if a.tipo in ("cpf", "email")} == {"confirmado"}


def test_arquivo_mudou_depois_da_previa(client, admin, gerente, ambiente, monkeypatch):
    achados = _varrer_e_confirmar(client, admin, gerente)
    c = _previa(client, gerente, achados)
    ambiente["servido"]["conteudo"] = csv_alunos() + b"\nOutro;1;2;3;4"
    _escrita(monkeypatch, True)
    r = client.post(f"/api/v1/lgpd/correcoes/{c['id']}/aprovar", headers=gerente, json={})
    assert r.status_code == 409 and "mudou" in r.json()["detail"] and not ambiente["uploads"]
    with SessionLocal() as db:
        cor = db.get(CorrecaoPublicada, c["id"])
        assert cor.status == "descartada" and cor.caminho_backup is None
    assert not any(ambiente["pasta"].rglob("*.csv"))   # arquivos com dado pessoal apagados


def test_validacoes_e_permissoes(client, admin, gerente, analista, ambiente):
    achados = _varrer_e_confirmar(client, admin, gerente)
    corpo = {"recurso_id": REC_CSV, "itens": [
        {"achado_id": achados["telefone"], "operador": "substituir"}]}
    r = client.post("/api/v1/lgpd/correcoes", headers=gerente, json=corpo)
    assert r.status_code == 422 and "confirmado" in r.json()["detail"]   # telefone não confirmado
    assert client.post("/api/v1/lgpd/correcoes", headers=analista, json=corpo).status_code == 403
    with SessionLocal() as db:
        db.get(Recurso, REC_CSV).url_type = None
        db.commit()
    corpo["itens"][0]["achado_id"] = achados["cpf"]
    r = client.post("/api/v1/lgpd/correcoes", headers=gerente, json=corpo)
    assert r.status_code == 422 and "link externo" in r.json()["detail"]
    lista = client.get("/api/v1/lgpd/correcoes", headers=analista).json()
    assert lista["escrita_habilitada"] is False and lista["operadores"][0]["valor"] == "mascara_cpf"
