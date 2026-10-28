"""Anonimização de recursos (US26) — funções puras (sem banco, sem FastAPI, sem rede).

Operadores (o do detector3 + a regra da GEDA para CPF):
  mascara_cpf   regra informada pela GEDA (17/09): oculta os 3 primeiros e os 2 últimos dígitos,
                mantendo a pontuação: 529.982.247-25 → ***.982.247-**
  substituir    troca pelo marcador do tipo: <CPF>, <EMAIL>, <ENDERECO>...
  mascarar      mantém os 3 primeiros e os 2 últimos caracteres (detector3); até 6 caracteres, tudo *
  hash          SHA-256 com sal secreto (12 caracteres): mesma pessoa, mesmo código, sem revelar o valor
  remover       apaga o trecho

Alcance de cada item:
  detectados    só os trechos que o motor encontra de novo na hora da prévia (célula pode ter texto
                além do dado: "cel (62) 99876-5432" → "cel <TELEFONE>")
  coluna        a célula inteira, em toda linha não vazia da coluna (inclusive valores que o motor não
                reconheceu, como um CPF com dígito errado)

A reescrita preserva o formato do arquivo: CSV com o mesmo encoding (BOM incluído), separador e
quebra de linha; XLSX com openpyxl, alterando só as células anonimizadas.
"""
from __future__ import annotations

import csv
import hashlib
import io
from collections.abc import Callable
from dataclasses import dataclass, field

from app.modules.lgpd.deteccao import Resultado

OPERADORES: dict[str, str] = {
    "mascara_cpf": "Ocultar 3 primeiros e 2 últimos dígitos (regra GEDA para CPF)",
    "substituir": "Substituir pelo tipo (<CPF>, <EMAIL>…)",
    "mascarar": "Mascarar (mantém 3 primeiros e 2 últimos caracteres)",
    "hash": "Código irreversível (hash com sal)",
    "remover": "Remover o valor",
}
ALCANCES: dict[str, str] = {
    "detectados": "Só os trechos detectados",
    "coluna": "A célula inteira, em toda a coluna",
}


def operador_padrao(tipo: str) -> str:
    """CPF segue a regra da GEDA; os demais tipos, o marcador do tipo."""
    return "mascara_cpf" if tipo == "cpf" else "substituir"


def mascara_cpf(valor: str) -> str:
    """Oculta o 1º, 2º, 3º, 10º e 11º dígitos, mantendo pontuação. Sem 11 dígitos → `mascarar`."""
    posicoes = [i for i, c in enumerate(valor) if c.isdigit()]
    if len(posicoes) != 11:
        return mascarar(valor)
    ocultar = set(posicoes[:3] + posicoes[-2:])
    return "".join("*" if i in ocultar else c for i, c in enumerate(valor))


def mascarar(valor: str) -> str:
    return "*" * len(valor) if len(valor) <= 6 else valor[:3] + "*" * (len(valor) - 5) + valor[-2:]


def aplicar(operador: str, valor: str, tipo: str, sal: str = "") -> str:
    if operador == "mascara_cpf":
        return mascara_cpf(valor)
    if operador == "substituir":
        return f"<{tipo.upper()}>"
    if operador == "mascarar":
        return mascarar(valor)
    if operador == "hash":
        return hashlib.sha256((sal + valor.strip()).encode("utf-8")).hexdigest()[:12]
    if operador == "remover":
        return ""
    raise ValueError(f"Operador desconhecido: {operador}")


def anonimizar_trechos(texto: str, trechos: list[tuple[Resultado, str, str]], sal: str = "") -> str:
    """Aplica operadores aos trechos [(resultado, tipo, operador)], do fim para o começo."""
    for r, tipo, op in sorted(trechos, key=lambda x: x[0].inicio, reverse=True):
        texto = texto[:r.inicio] + aplicar(op, texto[r.inicio:r.fim], tipo, sal) + texto[r.fim:]
    return texto


# ------------------------------------------------------------------ reescrita do arquivo
@dataclass
class Alteracoes:
    """(tabela, linha, índice da coluna) → novo valor. Tabela = aba do .xlsx ("" no CSV)."""

    celulas: dict[tuple[str, int, int], str] = field(default_factory=dict)

    def __len__(self) -> int:
        return len(self.celulas)


@dataclass
class FormatoCsv:
    encoding: str
    bom: bool
    separador: str
    quebra: str
    aspas_em_tudo: bool  # o arquivo original põe aspas em todos os campos (mantido na reescrita)


def formato_csv(conteudo: bytes) -> FormatoCsv:
    bom = conteudo.startswith(b"\xef\xbb\xbf")
    try:
        texto = conteudo.decode("utf-8-sig")
        encoding = "utf-8"
    except UnicodeDecodeError:
        texto, encoding = conteudo.decode("latin-1"), "latin-1"
    try:
        sep = csv.Sniffer().sniff(texto[:65536], delimiters=",;\t|").delimiter
    except csv.Error:
        sep = ","
    quebra = "\r\n" if "\r\n" in texto[:65536] else "\n"
    cabecalho = texto.split("\n", 1)[0].rstrip("\r")
    campos = cabecalho.split(sep)
    aspas = bool(cabecalho) and all(c.startswith('"') and c.endswith('"') and len(c) >= 2
                                    for c in campos)
    return FormatoCsv(encoding, bom, sep, quebra, aspas)


def reescrever_csv(conteudo: bytes, alteracoes: Alteracoes) -> bytes:
    """Mesma numeração de linhas de `leitura.ler_csv` (cabeçalho = 1, registros a partir de 2)."""
    f = formato_csv(conteudo)
    texto = conteudo.decode("utf-8-sig" if f.encoding == "utf-8" else f.encoding)
    saida = io.StringIO(newline="")
    escritor = csv.writer(saida, delimiter=f.separador, lineterminator=f.quebra,
                          quoting=csv.QUOTE_ALL if f.aspas_em_tudo else csv.QUOTE_MINIMAL)
    for n, row in enumerate(csv.reader(io.StringIO(texto, newline=""), delimiter=f.separador),
                            start=1):
        novo = list(row)
        for i in range(len(novo)):
            chave = ("", n, i)
            if chave in alteracoes.celulas:
                novo[i] = alteracoes.celulas[chave]
        escritor.writerow(novo)
    corpo = saida.getvalue()
    if not texto.endswith(("\n", "\r")):
        corpo = corpo.rstrip("\r\n")
    dados = corpo.encode(f.encoding, errors="replace")
    return (b"\xef\xbb\xbf" + dados) if f.bom else dados


def reescrever_xlsx(conteudo: bytes, alteracoes: Alteracoes) -> bytes:
    """Aba "" = planilha de uma aba só (como em `leitura.ler_xlsx`)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(conteudo))
    for (aba, linha, col), valor in alteracoes.celulas.items():
        ws = wb[aba] if aba else wb.worksheets[0]
        ws.cell(row=linha, column=col + 1).value = valor
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def reescrever(conteudo: bytes, formato: str, alteracoes: Alteracoes) -> bytes:
    if formato in ("CSV", "TSV", "TXT"):
        return reescrever_csv(conteudo, alteracoes)
    if formato == "XLSX":
        return reescrever_xlsx(conteudo, alteracoes)
    raise ValueError(f"Formato {formato} não pode ser reescrito.")


# ------------------------------------------------------------------ plano de alterações
@dataclass
class ItemPlano:
    """Um achado confirmado e como anonimizá-lo."""

    chave: str          # identificação livre (ex.: id do achado)
    tabela: str         # aba ("" no CSV)
    indice_coluna: int
    tipo: str           # código do sistema (cpf, email...)
    tipo_motor: str     # tipo do motor (CPF, EMAIL...)
    operador: str
    alcance: str
    alteradas: int = 0
    linhas_exemplo: list[int] = field(default_factory=list)


def planejar(item_coluna: dict[int, list[ItemPlano]], blocos, analisar: Callable, sal: str,
             alteracoes: Alteracoes, tabela: str) -> None:
    """Percorre os blocos de UMA tabela e preenche `alteracoes` para os itens das suas colunas.

    `item_coluna` = índice da coluna → itens (tipos) daquela coluna; `analisar(valores)` recebe
    [(linha, texto)] e devolve [(linha, texto, atipico, resultados)] (motor do detector3)."""
    for bloco in blocos:
        for i, itens in item_coluna.items():
            valores = [(n, row[i]) for n, row in bloco if i < len(row)]
            por_linha = {}
            if any(it.alcance == "detectados" for it in itens):
                por_linha = {n: rs for n, _v, _a, rs in analisar(valores, i)}
            for n, valor in valores:
                if not valor or not valor.strip():
                    continue
                inteira = next((it for it in itens if it.alcance == "coluna"), None)
                if inteira is not None:
                    novo = aplicar(inteira.operador, valor, inteira.tipo, sal)
                    _registrar(alteracoes, [inteira], tabela, n, i, novo)
                    continue
                trechos, afetados = [], []
                for r in por_linha.get(n, []):
                    for it in itens:
                        if it.tipo_motor == r.tipo:
                            trechos.append((r, it.tipo, it.operador))
                            if it not in afetados:
                                afetados.append(it)
                if trechos:
                    novo = anonimizar_trechos(valor, trechos, sal)
                    if novo != valor:
                        _registrar(alteracoes, afetados, tabela, n, i, novo)


def _registrar(alteracoes: Alteracoes, itens: list[ItemPlano], tabela: str, n: int, i: int,
               novo: str) -> None:
    alteracoes.celulas[(tabela, n, i)] = novo
    for it in itens:
        it.alteradas += 1
        if len(it.linhas_exemplo) < 3:
            it.linhas_exemplo.append(n)
