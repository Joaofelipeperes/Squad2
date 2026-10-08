"""Regras puras do módulo pda (sem banco, sem FastAPI) — reutilizáveis numa extensão CKAN.

- leitura da planilha do PDA (.xlsx/.csv) → linhas estruturadas (PBI-12);
- normalização da periodicidade prevista;
- extração do `name` a partir da URL "Disponível no Portal" (só para RESOLVER o ID — PBI-13);
- casamento da sigla do órgão com as organizações do portal;
- situação da base quanto ao prazo de abertura (PBI-19) e flag de prazo (PBI-20/22).
"""
import csv
import io
import re
import unicodedata
import zipfile
from collections.abc import Iterable, Iterator, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from itertools import chain, islice
from urllib.parse import unquote, urlsplit

# ------------------------------------------------------------------ texto


def sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", texto) if not unicodedata.combining(c))


def normalizar_texto(texto: object) -> str:
    """Sem acento, minúsculo, espaços colapsados, sem pontuação final ('Sigiloso?' → 'sigiloso')."""
    if texto is None:
        return ""
    t = sem_acento(str(texto)).lower()
    t = re.sub(r"\s+", " ", t).strip()
    return re.sub(r"[\s.:;,?!*]+$", "", t)


def _compacto(texto: str) -> str:
    """Só letras e dígitos (sem acento, minúsculo): 'GOIÁS TELECOM' → 'goiastelecom'."""
    return re.sub(r"[^a-z0-9]", "", normalizar_texto(texto))


# ------------------------------------------------------------------ periodicidade

SEM_INFORMACAO = "Sem informação"
MULTIPLA = "Múltipla"
DOMINIO_PERIODICIDADE: tuple[str, ...] = (
    "Diária", "Semanal", "Quinzenal", "Mensal", "Bimestral", "Trimestral", "Quadrimestral",
    "Semestral", "Anual", "Bienal", "Quadrienal", "Quinquenal", "Eventual", "Estática",
    MULTIPLA, SEM_INFORMACAO,
)

# Chave = texto normalizado (sem acento, minúsculo). Erros de digitação conhecidos da planilha
# GEDA entram explicitamente: não há correção "por semelhança" para não chutar periodicidade.
_PERIODICIDADES: dict[str, str] = {
    "diario": "Diária", "diaria": "Diária", "diariamente": "Diária",
    "semanal": "Semanal", "semanalmente": "Semanal",
    "quinzenal": "Quinzenal",
    "mensal": "Mensal", "mensalmente": "Mensal", "mnsal": "Mensal",
    "bimestral": "Bimestral",
    "trimestral": "Trimestral",
    "quadrimestral": "Quadrimestral",
    "semestral": "Semestral", "semstral": "Semestral",
    "anual": "Anual", "anualmente": "Anual",
    "bienal": "Bienal",
    # AMBÍGUO: "bianual" pode significar "a cada dois anos" (bienal) ou "duas vezes ao ano"
    # (semestral). Adotado Bienal conforme a especificação do Eixo 1 — pendente de validação GEDA.
    "bianual": "Bienal",
    "quadrienal": "Quadrienal",
    "quinquenal": "Quinquenal",
    "eventual": "Eventual", "sob demanda": "Eventual", "quando houver": "Eventual",
    "estatica": "Estática", "estatico": "Estática", "nao ha mais atualizacao": "Estática",
}
_SEM_INFO = {"", "n/a", "na", "n.a", "-", "--", "nao se aplica", "nao informado"}
# Ruído de digitação nas bordas da periodicidade ('.Semestral', '-Mensal', 'Mensal-')
_BORDAS_PERIODICIDADE = re.compile(r"^[\s.:;,?!*\-]+|[\s.:;,?!*\-]+$")


def _chave_periodicidade(texto: object) -> str:
    return _BORDAS_PERIODICIDADE.sub("", normalizar_texto(texto))


def normalizar_periodicidade(texto: object) -> str:
    """Periodicidade prevista da planilha → valor do domínio. O texto original é guardado à parte.

    Desconhecido → "Sem informação" (nunca chuta). Pontuação nas bordas é ignorada
    ('.Semestral', 'Mensal-'). Texto com "/": cada parte é normalizada; partes vazias ou N/A são
    ignoradas; se alguma parte for desconhecida → "Sem informação" ('Mensal/Quando necessário');
    duas ou mais periodicidades distintas → "Múltipla" ('Mensal/Trimestral/Semestral'); uma só →
    ela ('Mensal/Mnsal', 'Mensal/N/A' → Mensal).
    """
    chave = _chave_periodicidade(texto)
    if chave in _SEM_INFO:
        return SEM_INFORMACAO
    if chave in _PERIODICIDADES:
        return _PERIODICIDADES[chave]
    if "/" in chave:
        sem_na = re.sub(r"(?<![a-z0-9])n/a(?![a-z0-9])", "", chave)  # "Mensal/N/A" → "mensal/"
        partes = [_chave_periodicidade(p) for p in sem_na.split("/")]
        partes = [p for p in partes if p not in _SEM_INFO]
        if partes and all(p in _PERIODICIDADES for p in partes):
            distintas = {_PERIODICIDADES[p] for p in partes}
            return MULTIPLA if len(distintas) > 1 else distintas.pop()
    return SEM_INFORMACAO


# ------------------------------------------------------------------ vínculo pelo portal


def extrair_name_da_url(valor: object) -> str | None:
    """'https://dadosabertos.go.gov.br/dataset/<name>' → '<name>' (só para resolver o ID; o vínculo
    gravado é SEMPRE o ID do dataset — PBI-13). "Não", vazio ou URL sem /dataset/ → None.

    Ignora espaços, barra final, querystring, fragmento e maiúsculas (no CKAN o name é minúsculo).
    """
    if not isinstance(valor, str):
        return None
    texto = valor.strip()
    if not texto or "/dataset/" not in texto.lower():
        return None
    try:
        caminho = urlsplit(texto if "://" in texto else f"//{texto}").path
    except ValueError:
        return None
    segmentos = [unquote(s).strip() for s in caminho.split("/")]
    segmentos = [s for s in segmentos if s]
    for i, seg in enumerate(segmentos[:-1]):
        if seg.lower() == "dataset":
            name = segmentos[i + 1].lower()
            return name if re.fullmatch(r"[a-z0-9_\-]+", name) else None
    return None


# ------------------------------------------------------------------ órgão

_STOPWORDS = {"de", "da", "do", "das", "dos", "e", "a", "o"}
Orgao = tuple[str, str, str, str | None]  # (ckan_id, name, titulo, sigla)


def _palavras(texto: str) -> list[str]:
    return [p for p in re.split(r"[^a-z0-9]+", normalizar_texto(texto)) if p]


def _contem_sequencia(seq: list[str], tokens: list[str]) -> bool:
    """`seq` aparece como sequência CONTÍGUA de `tokens` (['casa', 'civil'] em
    ['secretaria', 'de', 'estado', 'da', 'casa', 'civil'])."""
    n = len(seq)
    return n > 0 and any(tokens[i:i + n] == seq for i in range(len(tokens) - n + 1))


def casar_orgao(sigla: object, orgaos: Iterable[Orgao]) -> str | None:
    """Sigla da planilha → ckan_id da organização, ou None se não houver correspondência ÚNICA.

    Critérios, em ordem (o primeiro com resultado decide; empate → None, nunca chuta):
      1) sigla igual a Organizacao.sigla;
      2) sigla compacta igual ao name inteiro sem hífens ou ao título inteiro compactado;
      3) sigla de UMA palavra como token inteiro do name (hífens) ou do título (espaço, hífen,
         parênteses, travessão); sigla de VÁRIAS palavras ("CASA CIVIL", "GOIÁS PARCERIAS") como
         sequência contígua desses tokens;
      4) sigla igual às iniciais das palavras significativas do título.
    """
    alvo = _compacto(sigla)
    if not alvo:
        return None
    palavras_sigla = _palavras(str(sigla))
    orgaos = list(orgaos)

    def criterio_tokens(o: Orgao) -> bool:
        toks_name, toks_titulo = _palavras(o[1]), _palavras(o[2])
        if alvo in toks_name or alvo in toks_titulo:  # sigla de uma palavra (comportamento original)
            return True
        return len(palavras_sigla) > 1 and (_contem_sequencia(palavras_sigla, toks_name)
                                            or _contem_sequencia(palavras_sigla, toks_titulo))

    criterios = (
        lambda o: bool(o[3]) and _compacto(o[3]) == alvo,
        lambda o: alvo in (_compacto(o[1]), _compacto(o[2])),
        criterio_tokens,
        lambda o: alvo == "".join(p[0] for p in _palavras(o[2]) if p not in _STOPWORDS),
    )
    for criterio in criterios:
        achados = {o[0] for o in orgaos if criterio(o)}
        if len(achados) == 1:
            return achados.pop()
        if len(achados) > 1:
            return None
    return None


def chave_sigla(sigla: str | None) -> str:
    """Sigla como chave de agrupamento: espaços colapsados, sem bordas, maiúscula ('Sead ' → 'SEAD')."""
    return re.sub(r"\s+", " ", sigla or "").strip().upper()


def orgao_por_evidencia(pares: Iterable[tuple[str, str | None]]) -> dict[str, str]:
    """Evidência do vínculo por ID: pares (sigla da planilha, organização do dataset vinculado) →
    {chave_sigla: organização} só para as siglas cujas bases vinculadas apontam para UMA única
    organização. Sigla com vínculos para duas ou mais organizações fica de fora (não chuta)."""
    evidencias: dict[str, set[str]] = {}
    for sigla, organizacao_id in pares:
        if organizacao_id:
            evidencias.setdefault(chave_sigla(sigla), set()).add(organizacao_id)
    return {s: next(iter(orgs)) for s, orgs in evidencias.items() if len(orgs) == 1}


# ------------------------------------------------------------------ situação e prazo

PUBLICADO = "Publicado"
SEM_RECURSO = "Sem recurso"
EM_ATRASO = "Em atraso"
PROXIMO_DO_PRAZO = "Próximo do prazo"
EM_DIA = "Em dia"
NAO_PUBLICADO = "Não publicado"
SITUACOES: tuple[str, ...] = (PUBLICADO, EM_DIA, EM_ATRASO, PROXIMO_DO_PRAZO, NAO_PUBLICADO,
                              SEM_RECURSO)


def situacao_base(*, vinculada: bool, dataset_ativo: bool | None, recursos_validos: int,
                  prazo: date | None, hoje: date, janela_dias: int) -> str:
    """PBI-19 — situação da base prevista quanto ao PRAZO DE ABERTURA (ressalva CGE-GO 17/09).

    `recursos_validos` exclui dicionário de dados. Dataset inativo no portal (excluído/privado,
    PBI-15) conta como não publicado.
    """
    if vinculada and dataset_ativo:
        return PUBLICADO if recursos_validos >= 1 else SEM_RECURSO
    if prazo is None:
        return NAO_PUBLICADO
    if prazo < hoje:
        return EM_ATRASO
    if prazo <= hoje + timedelta(days=janela_dias):
        return PROXIMO_DO_PRAZO
    return EM_DIA


def flag_prazo(prazo: date | None, hoje: date, janela_dias: int) -> str:
    """'vencido' | 'proximo' | '' — só pelo prazo, como pdaSituacaoDeadlineFlag (PBI-20/22)."""
    if prazo is None:
        return ""
    if prazo < hoje:
        return "vencido"
    if prazo <= hoje + timedelta(days=janela_dias):
        return "proximo"
    return ""


# ------------------------------------------------------------------ leitura da planilha

TAMANHO_MAXIMO = 5 * 1024 * 1024  # 5 MB (arquivo enviado)
EXTENSOES = (".xlsx", ".csv")
# Tetos contra planilha malformada ou "bomba" de compressão (o .xlsx é um zip): o arquivo de 5 MB
# pode descompactar em gigabytes e o openpyxl gera uma linha para cada número de linha saltado.
TAMANHO_DESCOMPACTADO_MAXIMO = 50 * 1024 * 1024  # soma dos arquivos dentro do .xlsx
ABAS_MAXIMO = 20
LINHAS_MAXIMO = 20_000           # linhas lidas da aba escolhida (incluindo as vazias)
COLUNAS_MAXIMO = 60              # a planilha do PDA tem ~10 colunas
VAZIAS_CONSECUTIVAS_MAXIMO = 1_000  # após isso a aba é dada como terminada

# cabeçalho normalizado → campo
_CABECALHOS: dict[str, str] = {
    "orgao": "orgao",
    "base de dados": "base",
    "descricao": "descricao",
    "unidade responsavel": "unidade",
    "atualizacao": "periodicidade",
    "periodicidade": "periodicidade",
    "politicas publicas": "politicas",
    "possui conteudo sigiloso": "sigiloso",
    "disponivel no portal": "portal",
    "prazo": "prazo",
    "prazo de abertura": "prazo",
}
_OBRIGATORIOS = {"orgao": "Órgão", "base": "Base de Dados"}
# Colunas opcionais (campo → rótulo mostrado ao usuário quando a coluna não é encontrada)
COLUNAS_OPCIONAIS: dict[str, str] = {
    "descricao": "Descrição", "unidade": "Unidade Responsável", "periodicidade": "Atualização",
    "politicas": "Políticas Públicas", "sigiloso": "Possui Conteúdo Sigiloso?",
    "portal": "Disponível no Portal", "prazo": "Prazo",
}
_LINHAS_PROCURA_CABECALHO = 30
_RODAPE = re.compile(r"^(total\b|nenhum filtro aplicado)")
_MSG_XLSX_INVALIDO = "Arquivo .xlsx ilegível ou corrompido."


class PlanilhaInvalida(ValueError):
    """Arquivo ilegível ou fora do formato esperado — a importação inteira é recusada."""


@dataclass
class LinhaPda:
    linha: int
    orgao_sigla: str
    nome_previsto: str
    descricao: str | None = None
    unidade_responsavel: str | None = None
    periodicidade_original: str | None = None
    politicas_publicas: str | None = None
    possui_conteudo_sigiloso: bool | None = None
    disponivel_no_portal: str | None = None
    prazo: date | None = None


@dataclass
class PlanilhaLida:
    linhas: list[LinhaPda] = field(default_factory=list)
    ignoradas: list[tuple[int, str]] = field(default_factory=list)
    colunas: frozenset[str] = frozenset()  # campos encontrados no cabeçalho

    @property
    def tem_coluna_prazo(self) -> bool:
        return "prazo" in self.colunas

    @property
    def colunas_opcionais_ausentes(self) -> list[str]:
        """Rótulos das colunas opcionais que a planilha não tem (aviso no resumo da importação)."""
        return [rotulo for campo, rotulo in COLUNAS_OPCIONAIS.items() if campo not in self.colunas]


def extensao_valida(nome_arquivo: str | None) -> str | None:
    nome = (nome_arquivo or "").strip().lower()
    return next((e for e in EXTENSOES if nome.endswith(e)), None)


def _texto(valor: object, colapsar: bool = True) -> str | None:
    """Célula → texto aparado (None se vazia). `colapsar=False` preserva quebras de linha."""
    if valor is None:
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    t = str(valor)
    t = re.sub(r"\s+", " ", t) if colapsar else t.replace("\r\n", "\n")
    return t.strip() or None


def interpretar_sim_nao(valor: object) -> bool | None:
    t = normalizar_texto(valor)
    if t in {"sim", "s", "yes", "x", "true", "1"}:
        return True
    if t in {"nao", "n", "no", "false", "0"}:
        return False
    return None


def interpretar_data(valor: object) -> date | None:
    """Prazo da planilha: date/datetime do xlsx, 'dd/mm/aaaa' ou 'aaaa-mm-dd' (com hora opcional,
    ignorada) ou serial do Excel. Vazio → None. Valor presente e ilegível → ValueError (nunca
    inventar prazo)."""
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, int | float) and not isinstance(valor, bool):
        if 36526 <= valor < 73051:  # serial do Excel entre 2000 e 2099 (ano solto não é data)
            return date(1899, 12, 30) + timedelta(days=int(valor))
        raise ValueError(str(valor))
    t = str(valor).strip()
    if not t:
        return None
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})(?:[ T]\d{1,2}:\d{2}(?::\d{2}(?:\.\d+)?)?)?", t)
    if m:
        return date(int(m[3]), int(m[2]), int(m[1]))
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T][\d:.]+)?", t)
    if m:
        return date(int(m[1]), int(m[2]), int(m[3]))
    raise ValueError(t)


def _decodificar(conteudo: bytes) -> str:
    """UTF-8 (com ou sem BOM); senão Windows-1252 (CSV salvo pelo Excel em pt-BR, superconjunto
    imprimível do Latin-1); por fim Latin-1, que decodifica qualquer byte."""
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            pass
    return conteudo.decode("latin-1")


Aba = Iterable[Sequence[object]]  # linhas de uma aba (ou de uma leitura do CSV), sob demanda


def _abas_csv(conteudo: bytes) -> Iterator[Aba]:
    """Uma leitura por separador candidato (";" e ","), a mais provável primeiro; vale a
    primeira em que o cabeçalho for encontrado (título acima do cabeçalho não atrapalha)."""
    texto = _decodificar(conteudo)
    primeira = next((ln for ln in texto.splitlines() if ln.strip()), "")
    separadores = [";", ","] if primeira.count(";") >= primeira.count(",") else [",", ";"]
    for sep in separadores:
        yield csv.reader(io.StringIO(texto, newline=""), delimiter=sep)


def _verificar_zip(conteudo: bytes) -> None:
    """Recusa .xlsx cujo conteúdo descompactado passa do teto, antes de o openpyxl abrir."""
    try:
        with zipfile.ZipFile(io.BytesIO(conteudo)) as zf:
            descompactado = sum(i.file_size for i in zf.infolist())
    except Exception as exc:  # não é um zip legível = .xlsx inválido (422)
        raise PlanilhaInvalida(_MSG_XLSX_INVALIDO) from exc
    if descompactado > TAMANHO_DESCOMPACTADO_MAXIMO:
        raise PlanilhaInvalida(
            f"Planilha .xlsx grande demais depois de descompactada (mais de "
            f"{TAMANHO_DESCOMPACTADO_MAXIMO // (1024 * 1024)} MB).")


def _protegido(linhas: Iterable[Sequence[object]]) -> Iterator[Sequence[object]]:
    """Falha do openpyxl DURANTE a leitura das linhas → PlanilhaInvalida (422), sem mascarar erros
    do restante do código (exceções do consumidor não passam por este gerador)."""
    try:
        yield from linhas
    except PlanilhaInvalida:
        raise
    except Exception as exc:  # qualquer falha do openpyxl = arquivo .xlsx inválido
        raise PlanilhaInvalida(_MSG_XLSX_INVALIDO) from exc


def _abas_xlsx(wb) -> Iterator[Aba]:
    if len(wb.worksheets) > ABAS_MAXIMO:
        raise PlanilhaInvalida(f"Planilha com mais de {ABAS_MAXIMO} abas.")
    for ws in wb.worksheets:
        ws.reset_dimensions()  # alguns geradores gravam a dimensão errada (ex.: A1:A1)
        yield _protegido(ws.iter_rows(max_col=COLUNAS_MAXIMO, values_only=True))


def _linhas_limitadas(linhas: Aba) -> Iterator[list[object]]:
    """Linhas da aba com os tetos: no máximo COLUNAS_MAXIMO colunas (sem as vazias à direita),
    erro acima de LINHAS_MAXIMO linhas e fim da aba após VAZIAS_CONSECUTIVAS_MAXIMO linhas vazias
    seguidas. As linhas vazias continuam sendo entregues (a numeração da planilha é preservada)."""
    vazias = 0
    for numero, linha in enumerate(linhas, start=1):
        if numero > LINHAS_MAXIMO:
            raise PlanilhaInvalida(
                f"Planilha com mais de {LINHAS_MAXIMO} linhas: confira se é a planilha do PDA.")
        valores = list(islice(linha, COLUNAS_MAXIMO))
        while valores and _texto(valores[-1]) is None:
            valores.pop()
        if valores:
            vazias = 0
        else:
            vazias += 1
            if vazias >= VAZIAS_CONSECUTIVAS_MAXIMO:
                return
        yield valores


def _mapear_cabecalho(linha: Sequence[object]) -> dict[str, int]:
    colunas: dict[str, int] = {}
    for i, valor in enumerate(linha):
        campo = _CABECALHOS.get(normalizar_texto(valor))
        if campo and campo not in colunas:
            colunas[campo] = i
    return colunas


def _localizar_cabecalho(abas: Iterable[Aba]):
    """Primeira aba cujo cabeçalho (nas 30 primeiras linhas) tem os campos obrigatórios →
    (linhas seguintes ao cabeçalho, sob demanda; índice 0-based do cabeçalho; colunas). Só as
    30 primeiras linhas de cada aba são lidas até o cabeçalho aparecer."""
    parcial: dict[str, int] = {}
    for aba in abas:
        linhas = _linhas_limitadas(aba)
        inicio = list(islice(linhas, _LINHAS_PROCURA_CABECALHO))
        for idx, linha in enumerate(inicio):
            colunas = _mapear_cabecalho(linha)
            if all(c in colunas for c in _OBRIGATORIOS):
                return chain(inicio[idx + 1:], linhas), idx, colunas
            if len(colunas) > len(parcial):
                parcial = colunas
    faltando = ", ".join(f'"{r}"' for c, r in _OBRIGATORIOS.items() if c not in parcial)
    raise PlanilhaInvalida(
        f"Cabeçalho obrigatório ausente: {faltando}. A planilha do PDA precisa das colunas "
        '"Órgão" e "Base de Dados" (demais colunas são opcionais).')


def _ler_linhas(abas: Iterable[Aba]) -> PlanilhaLida:
    linhas, idx_cab, col = _localizar_cabecalho(abas)

    def celula(linha: Sequence[object], campo: str) -> object:
        i = col.get(campo)
        return linha[i] if i is not None and i < len(linha) else None

    lida = PlanilhaLida(colunas=frozenset(col))
    prazos_invalidos: list[str] = []
    for numero, linha in enumerate(linhas, start=idx_cab + 2):
        if all(_texto(v) is None for v in linha):
            continue
        if _mapear_cabecalho(linha) == col:  # planilhas concatenadas ou convertidas de PDF
            lida.ignoradas.append((numero, "Cabeçalho repetido."))
            continue
        # Rodapé só na coluna Órgão: uma base cujo nome começa com "Total" é uma base de verdade
        if _RODAPE.match(normalizar_texto(celula(linha, "orgao"))):
            continue
        orgao, base = _texto(celula(linha, "orgao")), _texto(celula(linha, "base"))
        if not orgao or not base:
            falta = " e ".join(r for c, r in _OBRIGATORIOS.items()
                               if not {"orgao": orgao, "base": base}[c])
            lida.ignoradas.append((numero, f"Sem {falta}."))
            continue
        try:
            prazo = interpretar_data(celula(linha, "prazo"))
        except ValueError:
            prazos_invalidos.append(f"linha {numero} ({_texto(celula(linha, 'prazo'))!r})")
            continue
        lida.linhas.append(LinhaPda(
            linha=numero, orgao_sigla=orgao, nome_previsto=base,
            descricao=_texto(celula(linha, "descricao"), colapsar=False),
            unidade_responsavel=_texto(celula(linha, "unidade")),
            periodicidade_original=_texto(celula(linha, "periodicidade")),
            politicas_publicas=_texto(celula(linha, "politicas")),
            possui_conteudo_sigiloso=interpretar_sim_nao(celula(linha, "sigiloso")),
            disponivel_no_portal=_texto(celula(linha, "portal")),
            prazo=prazo,
        ))
    if prazos_invalidos:
        lista = "; ".join(prazos_invalidos[:10]) + ("; …" if len(prazos_invalidos) > 10 else "")
        raise PlanilhaInvalida(
            "Prazo em formato não reconhecido (use dd/mm/aaaa ou aaaa-mm-dd, com hora opcional): "
            f"{lista}.")
    return lida


def ler_planilha(conteudo: bytes, nome_arquivo: str) -> PlanilhaLida:
    """PBI-12 — lê a planilha do PDA (uma linha por base prevista). Colunas localizadas pelo
    cabeçalho normalizado, não pela posição. `linha` = número da linha na planilha (1-based).
    A leitura é sob demanda e limitada (abas, colunas, linhas e tamanho descompactado)."""
    ext = extensao_valida(nome_arquivo)
    if ext is None:
        raise PlanilhaInvalida("Formato não suportado: envie a planilha do PDA em .xlsx ou .csv.")
    if ext == ".csv":
        try:
            return _ler_linhas(_abas_csv(conteudo))
        except csv.Error as exc:
            raise PlanilhaInvalida(f"CSV ilegível: {exc}") from exc

    from openpyxl import load_workbook

    _verificar_zip(conteudo)
    try:
        wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as exc:  # qualquer falha do openpyxl ao abrir = arquivo .xlsx inválido (422)
        raise PlanilhaInvalida(_MSG_XLSX_INVALIDO) from exc
    try:
        return _ler_linhas(_abas_xlsx(wb))
    finally:
        wb.close()
